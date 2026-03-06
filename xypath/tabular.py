from __future__ import annotations

import csv
import io
import zipfile
from pathlib import Path

import xlrd
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


class Cell:
    def __init__(self, value, properties=None):
        self.value = value
        self.properties = {} if properties is None else properties


class RowSet:
    def __init__(self, rows, name="table", sheet=None):
        self._rows = rows
        self.name = name
        self.sheet = sheet

    def __iter__(self):
        return iter(self._rows)


class TableSet:
    def __init__(self, tables):
        self.tables = list(tables)

    def __iter__(self):
        return iter(self.tables)

    def __getitem__(self, key):
        if isinstance(key, int):
            return self.tables[key]
        for table in self.tables:
            if table.name == key:
                return table
        raise KeyError(key)


def _load_csv_rows(data):
    text = io.TextIOWrapper(io.BytesIO(data), newline="", encoding="utf-8")
    reader = csv.reader(text)
    rows = []
    for row in reader:
        rows.append([Cell(value) for value in row])
    return rows


def _load_zip_tables(data):
    """Load a zip archive as a tableset.

    CSV members keep the historical generic table name ("table") for compatibility
    with existing loader behavior and tests.
    """
    tables = []
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        for member in archive.namelist():
            if member.endswith("/"):
                continue
            with archive.open(member) as raw_file:
                rows = _load_csv_rows(raw_file.read())
            table_name = (
                # Keep CSV entry names as "table" for historical compatibility.
                "table"
                if Path(member).suffix.lower() == ".csv"
                else Path(member).stem
            )
            tables.append(RowSet(rows=rows, name=table_name))
    return tables


def _load_xls_tables(data):
    workbook = xlrd.open_workbook(file_contents=data, formatting_info=True)
    tables = []
    for sheet in workbook.sheets():
        rows = []
        for row_index in range(sheet.nrows):
            row = []
            for column_index in range(sheet.ncols):
                cell = sheet.cell(row_index, column_index)
                properties = {}
                if 0 <= cell.xf_index < len(workbook.xf_list):
                    xf = workbook.xf_list[cell.xf_index]
                    if 0 <= xf.font_index < len(workbook.font_list):
                        font = workbook.font_list[xf.font_index]
                        properties["bold"] = bool(font.bold)
                        properties["font_name"] = font.name
                row.append(Cell(cell.value, properties=properties))
            rows.append(row)
        tables.append(RowSet(rows=rows, name=sheet.name, sheet=sheet))
    return tables


def _load_xlsx_tables(data):
    tables = []
    workbook = load_workbook(
        filename=io.BytesIO(data),
        data_only=True,
        read_only=True,
    )
    try:
        for sheet in workbook.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                rows.append([Cell(value) for value in row])
            tables.append(RowSet(rows=rows, name=sheet.title, sheet=sheet))
    finally:
        workbook.close()
    return tables


def any_tableset(file_object, extension=""):
    data = file_object.read()
    extension = extension.strip(".").lower()

    if extension == "zip":
        return TableSet(_load_zip_tables(data))
    if extension in {"xls"}:
        return TableSet(_load_xls_tables(data))
    if extension in {"xlsx"}:
        return TableSet(_load_xlsx_tables(data))
    if extension in {"csv"}:
        return TableSet([RowSet(_load_csv_rows(data))])

    try:
        return TableSet(_load_zip_tables(data))
    except zipfile.BadZipFile:
        pass

    try:
        return TableSet(_load_xls_tables(data))
    except xlrd.XLRDError:
        pass

    try:
        return TableSet(_load_xlsx_tables(data))
    except (InvalidFileException, zipfile.BadZipFile):
        pass

    return TableSet([RowSet(_load_csv_rows(data))])
